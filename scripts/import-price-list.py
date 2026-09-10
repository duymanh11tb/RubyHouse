import html
import json
import re
import sys
import time
import urllib.request
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(sys.argv[1])
OUTPUT = ROOT / "server" / "price-list-2026.json"
UPLOADS = ROOT / "uploads" / "catalog"

LOCATIONS = {
    "RB.VAN PHUC": "van-phuc",
    "RB-LIEU GIAI": "lieu-giai",
    "RB-PHAN KE BINH": "phan-ke-binh",
}

COMMON_AMENITIES = [
    "Wi-Fi tốc độ cao",
    "Truyền hình cáp quốc tế và Nhật Bản",
    "Dọn phòng 3 lần/tuần",
    "Máy giặt và máy sấy riêng",
    "Lễ tân và an ninh 24/7",
    "Nước nóng trung tâm",
]


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def slug(value):
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def parse_album_rows(sheet):
    albums = []
    for row in sheet.iter_rows(values_only=True):
        values = [clean(v) for v in row]
        url = next((v for v in values if isinstance(v, str) and "drive.google.com/drive/folders/" in v), None)
        if not url:
            continue
        description = next((v for v in values if isinstance(v, str) and v != url), "")
        codes = re.findall(r"R\d{3}", description.upper())
        folder_match = re.search(r"/folders/([^?/#]+)", url)
        if codes and folder_match:
            albums.append({"codes": codes, "url": url, "folder_id": folder_match.group(1)})
    return albums


def drive_files(folder_id):
    url = f"https://drive.google.com/drive/folders/{folder_id}"
    body = None
    for attempt in range(3):
        try:
            request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
            body = urllib.request.urlopen(request, timeout=90).read().decode("utf-8", "ignore")
            break
        except (TimeoutError, OSError):
            if attempt == 2:
                raise
            time.sleep(1)
    matches = re.findall(
        r'aria-label="([^"]+?\.(?:jpe?g|png|webp|gif|mp4|webm|mov))[^\"]*"[^>]*?data-handled-by-drag-and-drop.*?data-id="([^"]+)"',
        body,
        flags=re.IGNORECASE | re.DOTALL,
    )
    seen = set()
    files = []
    for raw_name, file_id in matches:
        if file_id in seen:
            continue
        seen.add(file_id)
        name = html.unescape(raw_name)
        ext = Path(name).suffix.lower()
        files.append({"id": file_id, "name": name, "kind": "video" if ext in {".mp4", ".webm", ".mov"} else "image"})
    return files


def download_image(file_id, destination):
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists() and destination.stat().st_size > 0:
        return
    for attempt in range(3):
        try:
            request = urllib.request.Request(
                f"https://lh3.googleusercontent.com/d/{file_id}=w1600",
                headers={"User-Agent": "Mozilla/5.0"},
            )
            with urllib.request.urlopen(request, timeout=90) as response:
                destination.write_bytes(response.read())
            return
        except (TimeoutError, OSError):
            if attempt == 2:
                raise
            time.sleep(1)


def parse_bedrooms(description):
    if re.search(r"\bstudio\b", description, re.IGNORECASE):
        return 0
    match = re.search(r"\b0?([1-9])\s*bedroom", description, re.IGNORECASE)
    return int(match.group(1)) if match else 0


def main():
    workbook = openpyxl.load_workbook(SOURCE, data_only=True)
    rooms = []
    media_cache = {}

    for sheet_name, location_id in LOCATIONS.items():
        sheet = workbook[sheet_name]
        albums = parse_album_rows(sheet)
        album_for_room = {code: album for album in albums for code in album["codes"]}

        for row in sheet.iter_rows(min_row=10, values_only=True):
            room_no = clean(row[0])
            if not isinstance(room_no, str) or not re.fullmatch(r"R\d{3}", room_no.upper()):
                continue
            room_no = room_no.upper()
            description = str(clean(row[1]) or "Căn hộ dịch vụ đầy đủ nội thất.")
            area_match = re.search(r"(\d+)\s*m2", description, re.IGNORECASE)
            area = int(area_match.group(1)) if area_match else 1
            album = album_for_room.get(room_no)
            media = []

            if album:
                cache_key = (location_id, album["folder_id"])
                if cache_key not in media_cache:
                    group_dir = UPLOADS / location_id / slug("-".join(album["codes"]))
                    imported = []
                    for index, item in enumerate(drive_files(album["folder_id"]), 1):
                        if item["kind"] != "image":
                            imported.append({
                                "kind": "video",
                                "url": f"https://drive.google.com/file/d/{item['id']}/view",
                                "alt": item["name"],
                            })
                            continue
                        destination = group_dir / f"{index:02d}.jpg"
                        download_image(item["id"], destination)
                        imported.append({
                            "kind": "image",
                            "url": "/" + destination.relative_to(ROOT).as_posix(),
                            "alt": f"Ruby House {location_id.replace('-', ' ').title()} – {room_no}",
                        })
                    media_cache[cache_key] = imported[:50]
                media = media_cache[cache_key]

            monthly = clean(row[2])
            promotion = clean(row[3])
            daily = clean(row[4])
            status = str(clean(row[5]) or "").lower()
            rooms.append({
                "id": f"{location_id}-{room_no.lower()}",
                "location_id": location_id,
                "name": f"Phòng {room_no}",
                "room_no": room_no,
                "area": area,
                "bedrooms": parse_bedrooms(description),
                "description": description,
                "amenities": COMMON_AMENITIES,
                "monthly_price": monthly,
                "promotion_price": promotion,
                "daily_price": daily,
                "availability": "available" if "avail" in status else "occupied",
                "album_url": album["url"] if album else None,
                "published": True,
                "media": media,
            })

    OUTPUT.write_text(json.dumps({"rooms": rooms}, ensure_ascii=False, indent=2), encoding="utf-8")
    images = sum(1 for values in media_cache.values() for item in values if item["kind"] == "image")
    videos = sum(1 for values in media_cache.values() for item in values if item["kind"] == "video")
    print(f"Imported {len(rooms)} rooms, {images} images, {videos} videos into {OUTPUT}")


if __name__ == "__main__":
    main()
